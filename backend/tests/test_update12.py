"""Update 12: routing words, agent trace/usage/evals, company hub extras, alerts, restore drill, time tracker, compare, quizzes, search."""
import json
import smtplib
from datetime import date, timedelta

import pytest

from app.services import agent_evals, agents, alerts, company, drill, llm_router, mailer, pm, providers, quizzes, search, state, usage
from app.services.providers import Call, Turn

KEYS = ("GOOGLE_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY")


@pytest.fixture()
def keys(monkeypatch):
    for k in KEYS:
        monkeypatch.setenv(k, "x")
    llm_router.reset_usage()


class Fake:
    provider, model = "anthropic", "fake-model"

    def __init__(self, turns):
        self.turns, self.i = turns, 0

    def start(self, system, tools, question):
        self.system, self.tools, self.question = system, tools, question

    def next_turn(self):
        t = self.turns[min(self.i, len(self.turns) - 1)]
        self.i += 1
        return t

    def add_results(self, turn, results):
        pass


def fakes(*turns):
    f = Fake(list(turns))
    return {"anthropic": f, "openai": f, "google": f}, f


# ------------------------------------------------------------------ routing words
@pytest.mark.parametrize("q", ["Analyze my project risks against my sprint progress and recommend what to do first", "Evaluate the trade-offs of two designs",
                               "Assess our delivery strategy", "Optimise the schedule", "Diagnose the margin drop", "Simulate three scenarios"])
def test_analysis_words_now_route_to_the_complex_tier(q):
    assert llm_router.classify(q)[0] == "complex"


@pytest.mark.parametrize("q,kind", [("Compare revenue by entity", "medium"), ("Summarise last week", "medium"), ("hello", "simple"), ("What is my top open risk?", "simple"),
                                    ("What does the analyst track cover?", "simple")])
def test_ordinary_questions_keep_their_tier(q, kind):
    assert llm_router.classify(q)[0] == kind


def test_every_eval_case_expects_the_tier_the_router_gives():
    for t, cases in agent_evals.CASES.items():
        assert t in agents.AGENTS
        for c in cases:
            assert llm_router.classify(c["ask"])[0] == c["tier"], (t, c["id"])


def test_plain_text_call_sends_no_tools(keys):
    ad, f = fakes(Turn("Brief text."))
    r = llm_router.complete("sys", "prompt", "medium", "auto", ad)
    assert r["text"] == "Brief text." and f.tools == [] and r["kind"] == "medium"
    with pytest.raises(Exception):
        llm_router.complete("s", "p", "medium", "auto", fakes(Turn(""))[0])


def test_adapters_leave_the_tools_field_out_when_there_are_none():
    class Anth:
        def __init__(self):
            self.kw = None
            self.messages = self

        def create(self, **kw):
            self.kw = kw
            blk = type("B", (), {"type": "text", "text": "hi"})()
            return type("R", (), {"content": [blk], "stop_reason": "end_turn", "usage": None})()
    c = Anth()
    a = providers.AnthropicAdapter(client=c, model="m")
    a.start("sys", [], "q")
    a.next_turn()
    assert "tools" not in c.kw


# ------------------------------------------------------------------ agent trace and usage
def test_tool_trace_says_what_each_call_did(client, keys):
    ad, _ = fakes(Turn("", [Call("1", "run_sql", {"sql": "SELECT COUNT(*) AS n FROM fact_operations"})]), Turn("Six records."))
    r = agents.chat("analyst", "how many records are there?", adapters=ad, provider="anthropic")
    t = r.tools_used[0]
    assert t["tool"] == "run_sql" and "COUNT(*)" in t["what"] and t["chars"] > 0 and r.reply == "Six records."


def test_usage_counts_tier_provider_and_track(client, keys):
    ad, _ = fakes(Turn("ok"))
    agents.chat("pm", "What's my top open risk?", adapters=ad)
    agents.chat("pm", "Why is my sprint behind, and what is the root cause?", adapters=ad)
    u = usage.summary(7)
    assert u["total_questions"] == 2 and u["by_track"] == {"pm": 2}
    assert sum(u["matrix"]["simple"].values()) == 1 and sum(u["matrix"]["complex"].values()) == 1
    assert u["total_cost_usd"] is None and "No prices set" in u["cost_note"]
    assert len(u["daily"]) == 7 and u["recent"][0]["track"] == "pm"
    assert client.get("/api/ai/usage?days=3").json()["days"] == 3


def test_cost_appears_only_when_you_set_prices(client, keys, monkeypatch):
    state.run("INSERT INTO ai_log (created_at, question, provider, model, kind, input_tokens, output_tokens, ok, charted, workspace, tier) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
              (state.now(), "q", "openai", "m", "agent", 1_000_000, 500_000, 1, 0, "practice", "medium"))
    assert usage.summary(3)["tokens"]["openai"]["cost_usd"] is None
    monkeypatch.setenv("BI_PRICE_OPENAI_IN", "2"); monkeypatch.setenv("BI_PRICE_OPENAI_OUT", "8")
    u = usage.summary(3)
    assert u["tokens"]["openai"]["cost_usd"] == 6.0 and u["total_cost_usd"] == 6.0


def test_old_state_file_gets_the_new_log_columns(tmp_path, monkeypatch):
    import sqlite3
    monkeypatch.setenv("BI_DB_PATH", str(tmp_path / "x.duckdb"))
    state.close_all()
    c = sqlite3.connect(tmp_path / "x_state.sqlite")
    c.execute("CREATE TABLE ai_log (id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT, question TEXT, provider TEXT, model TEXT, kind TEXT, input_tokens INTEGER, output_tokens INTEGER, ok INTEGER, charted INTEGER, workspace TEXT)")
    c.commit(); c.close()
    cols = {r[1] for r in state.conn().execute("PRAGMA table_info(ai_log)").fetchall()}
    state.close_all()
    assert {"tier", "track"} <= cols


# ------------------------------------------------------------------ evals
def test_route_eval_is_free_and_saved(client):
    r = client.post("/api/agent-evals/run", json={"track": "pm"}).json()
    assert r["mode"] == "route" and r["pct"] == 100 and r["total"] == len(agent_evals.CASES["pm"])
    ov = client.get("/api/agent-evals").json()
    assert ov["runs"][0]["track"] == "pm" and {t["track"] for t in ov["tracks"]} == set(agents.AGENTS)
    assert client.post("/api/agent-evals/run", json={"track": "zzz"}).status_code == 404


def test_live_eval_needs_confirmation_and_checks_structure(client, keys):
    assert client.post("/api/agent-evals/run", json={"track": "pm", "mode": "live"}).status_code == 400
    ad, _ = fakes(Turn("A reply with no tool use."))
    r = agent_evals.run("pm", "live", "auto", True, ad)
    assert r["mode"] == "live" and r["total"] == 4 and r["passed"] < r["total"]            # a model that never uses tools fails those checks
    risk = next(x for x in r["results"] if x["id"] == "risk")
    assert any(c["name"] == "used a data tool" and not c["ok"] for c in risk["checks"])
    assert r["models"] == ["anthropic:fake-model"] or r["models"] == ["openai:fake-model"] or r["models"]
    assert len(agent_evals.history()) == 1


def test_live_eval_stops_when_ai_is_off(client, keys):
    client.post("/api/workspaces/active", json={"name": "real"})            # Real defaults to AI off
    with pytest.raises(agent_evals.EvalError) as e:
        agent_evals.run("pm", "live", "auto", True, fakes(Turn("x"))[0])
    assert e.value.status == 403


# ------------------------------------------------------------------ company hub
def test_notes_and_due_dates_flow_into_the_plan(client):
    t = date.today()
    ov = client.put("/api/company/deliverables/req-capture/meta", json={"note": "ask Dana", "due": (t - timedelta(days=2)).isoformat()}).json()
    it = next(i for p in ov["phases"] for i in p["deliverables"] if i["id"] == "req-capture")
    assert it["note"] == "ask Dana" and it["overdue"] and ov["overdue"] == 1
    ov = client.put("/api/company/deliverables/pm-risks/meta", json={"due": (t + timedelta(days=3)).isoformat()}).json()
    assert ov["due_soon"] == 1
    assert client.put("/api/company/deliverables/pm-risks/meta", json={"due": "15/11/2026"}).status_code == 400
    assert client.put("/api/company/deliverables/pm-risks/meta", json={"due": "2026-02-31"}).status_code == 400
    assert client.put("/api/company/deliverables/nope/meta", json={"note": "x"}).status_code == 404
    ov = client.put("/api/company/deliverables/pm-risks/meta", json={"note": "", "due": ""}).json()
    assert ov["due_soon"] == 0


def test_finished_deliverables_are_never_overdue(client):
    client.put("/api/company/deliverables/sa-model/meta", json={"due": "2020-01-01"})
    ov = client.put("/api/company/deliverables/sa-model", json={"done": True}).json()
    assert ov["overdue"] == 0


def test_snapshot_diff_names_what_moved(client):
    first = client.post("/api/company/snapshots").json()
    assert "first snapshot" in first["brief"]
    client.put("/api/company/deliverables/sa-model", json={"done": True})
    client.put("/api/company/deliverables/fs-api", json={"done": True})
    second = client.post("/api/company/snapshots").json()
    assert "Document the data model" in second["brief"] and "Design the API" in second["brief"] and "+" in second["brief"]
    snaps = client.get("/api/company/snapshots").json()
    assert len(snaps["snapshots"]) == 2 and snaps["snapshots"][0]["done"] == second["done"] and snaps["weekly_due"] is False


def test_ai_brief_sees_only_counts_and_titles(client, keys):
    ad, f = fakes(Turn("Two deliverables moved. Next: capture requirements."))
    client.post("/api/company/snapshots")
    r = company.ai_brief(ad)
    assert r["ai"] and "Two deliverables" in r["brief"] and r["provider"]
    assert "fact_operations" not in f.question and "revenue" not in f.question.lower().replace("deliverable", "")
    assert company.snapshots(2)[0]["ai"] is True


def test_ai_brief_is_blocked_when_ai_is_off(client, keys):
    client.post("/api/workspaces/active", json={"name": "real"})
    assert client.post("/api/company/snapshots/ai-brief").status_code == 403


def test_company_exports_open_as_real_files(client):
    import io
    from openpyxl import load_workbook
    client.put("/api/company/deliverables/req-capture/meta", json={"note": "n1", "due": "2020-01-01"})
    client.post("/api/company/snapshots")
    x = client.get("/api/company/export?format=xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Deliverables", "Disciplines", "Snapshots"]
    rows = list(wb["Deliverables"].iter_rows(values_only=True))
    assert len(rows) == 23 and any(r[3] == "OVERDUE" for r in rows[1:])
    p = client.get("/api/company/export?format=pdf")
    assert p.content[:4] == b"%PDF" and p.headers["content-type"] == "application/pdf"
    assert client.get("/api/company/export?format=doc").status_code == 400


def test_weekly_job_snapshots_once_a_week(client):
    from app.services import jobs
    assert "snapshot" in jobs.tick()
    assert "snapshot" not in jobs.tick()
    state.run("UPDATE company_snapshots SET at = '2020-01-01 00:00:00'")
    assert company.weekly_due() is True


# ------------------------------------------------------------------ alerts and mail
def test_mailer_refuses_without_config_and_sends_to_listed_addresses_only(monkeypatch):
    for k in ("BI_SMTP_HOST", "BI_REPORT_TO", "BI_SMTP_USER", "BI_SMTP_FROM"):
        monkeypatch.delenv(k, raising=False)
    with pytest.raises(mailer.MailError) as e:
        mailer.send("s", "b")
    assert e.value.status == 503
    monkeypatch.setenv("BI_SMTP_HOST", "smtp.test"); monkeypatch.setenv("BI_REPORT_TO", "me@example.com, you@example.com"); monkeypatch.setenv("BI_SMTP_FROM", "bi@example.com")
    sent = []

    class S:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def send_message(self, m): sent.append(m)
    assert mailer.send("Hello", "Body", lambda c: S()) == ["me@example.com", "you@example.com"]
    assert sent[0]["To"] == "me@example.com, you@example.com" and sent[0]["Subject"] == "Hello"


def mail_env(monkeypatch):
    monkeypatch.setenv("BI_SMTP_HOST", "smtp.test"); monkeypatch.setenv("BI_REPORT_TO", "me@example.com"); monkeypatch.setenv("BI_SMTP_FROM", "bi@example.com")
    sent = []

    class S:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def send_message(self, m): sent.append(m)
    return sent, lambda c: S()


def test_alerts_see_red_kpis_and_missing_backup(client):
    client.post("/api/kpis", json={"name": "Impossible margin", "metric": "margin_pct", "direction": "higher", "target": 99, "warn_pct": 1, "window_days": 0})
    ids = {a["id"].split(":")[0] for a in alerts.current()}
    assert "kpi" in ids and "backup" in ids
    s = client.get("/api/alerts").json()
    assert s["red"] >= 2 and s["config"]["enabled"] is False and s["email_ready"] in (True, False)


def test_alert_config_validates_and_email_needs_smtp(client, monkeypatch):
    for k in ("BI_SMTP_HOST", "BI_REPORT_TO"):
        monkeypatch.delenv(k, raising=False)
    assert client.put("/api/alerts/config", json={"email": True}).status_code == 503
    assert client.put("/api/alerts/config", json={"stale_days": 0}).status_code == 400
    assert client.put("/api/alerts/config", json={"cooldown_hours": "x"}).status_code == 400
    ok = client.put("/api/alerts/config", json={"enabled": True, "stale_days": 3}).json()
    assert ok["config"]["enabled"] and ok["config"]["stale_days"] == 3
    assert client.post("/api/alerts/test-email").status_code == 503


def test_alert_email_is_sent_once_per_cooldown_and_again_after_it_clears(client, monkeypatch):
    sent, fac = mail_env(monkeypatch)
    client.put("/api/alerts/config", json={"enabled": True, "email": True})
    client.post("/api/kpis", json={"name": "K", "metric": "margin_pct", "direction": "higher", "target": 99, "warn_pct": 1, "window_days": 0})
    r1 = alerts.check(send=True, smtp_factory=fac)
    assert len(sent) == 1 and r1["emailed"] and "RED" in sent[0].get_content()
    alerts.check(send=True, smtp_factory=fac)
    assert len(sent) == 1                                                  # same alerts, inside the cooldown
    ws = client.get("/api/workspaces").json()["active"]
    kid = client.get("/api/kpis").json()["kpis"][0]["id"]
    client.delete(f"/api/kpis/{kid}")
    alerts.check(send=True, smtp_factory=fac)                              # the KPI alert cleared, so it is forgotten
    assert not any(k.startswith("kpi:") for k in json.loads(state.kv_get(f"alerts_sent:{ws}")))
    n = len(sent)
    client.post("/api/kpis", json={"name": "K", "metric": "margin_pct", "direction": "higher", "target": 99, "warn_pct": 1, "window_days": 0})
    alerts.check(send=True, smtp_factory=fac)                              # a new red KPI is a new alert: emailed
    assert len(sent) == n + 1


def test_a_failing_mail_server_is_reported_not_raised(client, monkeypatch):
    mail_env(monkeypatch)
    client.put("/api/alerts/config", json={"enabled": True, "email": True})
    client.post("/api/kpis", json={"name": "K", "metric": "margin_pct", "direction": "higher", "target": 99, "warn_pct": 1, "window_days": 0})

    def boom(c):
        raise smtplib.SMTPException("down")
    r = alerts.check(send=True, smtp_factory=boom)
    assert r["emailed"] == [] and "Couldn't send" in r["email_problem"]


def test_stale_data_is_only_flagged_in_real(client):
    assert not any(a["id"] == "data:stale" for a in alerts.current())
    client.post("/api/workspaces/active", json={"name": "real"})
    from app.services.db import get_cursor
    with get_cursor() as cur:
        cur.execute("INSERT INTO dim_entities VALUES ('E1','Old','Fleet',1.0)")
        cur.execute("INSERT INTO fact_operations (fact_id, record_date, entity_id, revenue, operational_cost, units_processed, duration_minutes, status) VALUES ('f1','2020-01-01','E1',1,1,1,1,'Completed')")
    assert any(a["id"] == "data:stale" for a in alerts.current())


# ------------------------------------------------------------------ restore drill
def test_drill_needs_a_backup_then_passes(client):
    assert client.post("/api/dba/drill").status_code == 404
    client.post("/api/backups", json={})
    d = client.post("/api/dba/drill").json()
    assert d["ok"] and d["reading"].startswith("PASSED") and all(s["ok"] for s in d["steps"]) and d["seconds"] >= 0
    h = client.get("/api/dba/drill").json()
    assert len(h["history"]) == 1 and h["due"] is False


def test_drill_catches_a_damaged_backup_and_leaves_live_data_alone(client, tmp_path):
    from app.config import backups_dir
    from app.services.db import get_cursor
    name = client.post("/api/backups", json={}).json()["name"]
    with get_cursor() as cur:
        before = cur.execute("SELECT COUNT(*) FROM fact_operations").fetchone()[0]
    f = next(backups_dir().rglob(name))
    f.write_bytes(b"this is not a database")
    d = client.post("/api/dba/drill", json={"name": name}).json()
    assert d["ok"] is False and d["reading"].startswith("FAILED")
    with get_cursor() as cur:
        assert cur.execute("SELECT COUNT(*) FROM fact_operations").fetchone()[0] == before
    assert any(a["id"] == "drill:failed" for a in alerts.current())


def test_drill_reports_how_far_behind_the_backup_is(client):
    from app.services.db import get_cursor
    client.post("/api/backups", json={})
    with get_cursor() as cur:
        cur.execute("INSERT INTO fact_operations (fact_id, record_date, entity_id, revenue, operational_cost, units_processed, duration_minutes, status) SELECT 'new1', current_date, entity_id, 5, 1, 1, 1, 'Completed' FROM dim_entities LIMIT 1")
    d = client.post("/api/dba/drill").json()
    assert d["ok"] and d["records_since_backup"] == 1 and "1 records" in d["recovery_point"]


# ------------------------------------------------------------------ time tracker
def test_logged_time_replaces_typed_cost_in_earned_value(client, monkeypatch):
    t = date(2026, 6, 15)
    monkeypatch.setattr(pm, "today", lambda: t)
    a = client.post("/api/pm/items", json={"title": "A", "status": "done", "planned_cost": 1000, "actual_cost": 900, "start_date": "2026-06-01", "duration_days": 5}).json()
    ev = client.get("/api/pm").json()["earned_value"]
    assert ev["ac"] == 900 and ev["items_costed_from_time"] == 0
    client.put("/api/pm/rate", json={"rate": 100})
    v = client.post("/api/pm/time", json={"item_id": a["id"], "day": "2026-06-10", "hours": 8, "note": "build"}).json()
    assert v["total_hours"] == 8 and v["cost"] == 800 and v["feeds_earned_value"]
    ev = client.get("/api/pm").json()["earned_value"]
    assert ev["ac"] == 800 and ev["cpi"] == 1.25 and ev["items_costed_from_time"] == 1


def test_time_without_a_rate_does_not_touch_earned_value(client, monkeypatch):
    monkeypatch.setattr(pm, "today", lambda: date(2026, 6, 15))
    a = client.post("/api/pm/items", json={"title": "A", "status": "done", "planned_cost": 1000, "actual_cost": 900}).json()
    v = client.post("/api/pm/time", json={"item_id": a["id"], "day": "2026-06-10", "hours": 3}).json()
    assert v["feeds_earned_value"] is False and v["cost"] is None
    assert client.get("/api/pm").json()["earned_value"]["ac"] == 900


@pytest.mark.parametrize("body,code", [({"hours": 0}, 400), ({"hours": 25}, 400), ({"hours": "x"}, 400), ({"hours": 1, "day": "2999-01-01"}, 400),
                                       ({"hours": 1, "day": "nope"}, 400), ({"hours": 1, "item_id": 9999}, 404), ({"hours": 1, "item_id": "abc"}, 400)])
def test_time_entries_are_validated(client, body, code):
    assert client.post("/api/pm/time", json=body).status_code == code


def test_time_summary_unassigned_delete_and_rate_limits(client):
    a = client.post("/api/pm/items", json={"title": "A", "planned_cost": 100}).json()
    client.put("/api/pm/rate", json={"rate": 50})
    client.post("/api/pm/time", json={"item_id": a["id"], "hours": 4})
    v = client.post("/api/pm/time", json={"hours": 2, "note": "admin"}).json()
    assert v["total_hours"] == 6 and v["unassigned_hours"] == 2 and v["by_item"][0]["over_plan"] is True and len(v["entries"]) == 2
    tid = v["entries"][0]["id"]
    assert client.delete(f"/api/pm/time/{tid}").json()["total_hours"] in (4, 2)
    assert client.delete("/api/pm/time/99999").status_code == 404
    assert client.put("/api/pm/rate", json={"rate": -1}).status_code == 400 and client.put("/api/pm/rate", json={"rate": "x"}).status_code == 400


def test_time_belongs_to_one_workspace(client):
    client.post("/api/pm/time", json={"hours": 2})
    client.post("/api/workspaces/active", json={"name": "real"})
    assert client.get("/api/pm/time").json()["total_hours"] == 0


# ------------------------------------------------------------------ compare
def test_compare_before_and_after_real_has_data(client):
    c = client.get("/api/compare").json()
    assert c["exists"]["practice"] and c["real_has_data"] is False and "nothing to compare" in c["note"]
    client.post("/api/workspaces/active", json={"name": "real"})
    from app.services.db import get_cursor
    with get_cursor() as cur:
        cur.execute("INSERT INTO dim_entities VALUES ('R1','Real One','Fleet',50.0)")
        cur.execute("INSERT INTO fact_operations (fact_id, record_date, entity_id, revenue, operational_cost, units_processed, duration_minutes, status) VALUES ('r1','2026-09-01','R1',200,50,4,30,'Completed')")
    client.post("/api/workspaces/active", json={"name": "practice"})
    c = client.get("/api/compare").json()
    assert c["real_has_data"] and not c["errors"]
    m = {x["key"]: x for x in c["metrics"]}
    assert m["records"]["real"] == 1 and m["revenue"]["real"] == 200 and m["margin_pct"]["real"] == 75.0 and m["records"]["practice"] > 1
    assert m["revenue"]["delta"]["abs"] == round(200 - m["revenue"]["practice"], 2)
    assert any(r["table"] == "fact_operations" for r in c["tables"])


def test_compare_never_locks_or_changes_the_files(client):
    from app.services.db import get_cursor
    client.post("/api/workspaces/active", json={"name": "real"})
    client.post("/api/workspaces/active", json={"name": "practice"})
    for _ in range(2):
        assert client.get("/api/compare").status_code == 200
    with get_cursor() as cur:
        assert cur.execute("SELECT COUNT(*) FROM fact_operations").fetchone()[0] > 0
    client.post("/api/workspaces/active", json={"name": "real"})
    assert client.get("/api/compare").status_code == 200


def test_same_kpi_definitions_are_run_on_both_datasets(client):
    client.post("/api/kpis/examples")
    client.post("/api/workspaces/active", json={"name": "real"})
    from app.services.db import get_cursor
    with get_cursor() as cur:
        cur.execute("INSERT INTO dim_entities VALUES ('R1','Real One','Fleet',50.0)")
        cur.execute("INSERT INTO fact_operations (fact_id, record_date, entity_id, revenue, operational_cost, units_processed, duration_minutes, status) VALUES ('r1','2026-09-01','R1',200,50,4,30,'Completed')")
    client.post("/api/workspaces/active", json={"name": "practice"})
    c = client.get("/api/compare").json()
    assert c["kpis"] and all(k["practice"] and k["real"] for k in c["kpis"])


# ------------------------------------------------------------------ quizzes
def test_quizzes_cover_every_track_with_valid_questions():
    from app.services import manuals
    assert set(quizzes.QUIZZES) == set(manuals.MANUALS)
    for t, qs in quizzes.QUIZZES.items():
        ids = {s["id"] for s in manuals.MANUALS[t]["steps"]}
        assert len(qs) >= 5
        for q in qs:
            assert q["step"] in ids and len(q["options"]) == 4 and 0 <= q["answer"] < 4 and len(set(q["options"])) == 4 and q["why"]
    assert len({q["answer"] for qs in quizzes.QUIZZES.values() for q in qs}) == 4          # the right answer isn't always in one slot


def test_the_questions_endpoint_never_leaks_the_answers(client):
    d = client.get("/api/quizzes/pm").json()
    assert "answer" not in json.dumps(d) and "why" not in json.dumps(d) and len(d["questions"]) == len(quizzes.QUIZZES["pm"])
    assert client.get("/api/quizzes/zzz").status_code == 404


def test_grading_scoring_and_best_score(client):
    qs = quizzes.QUIZZES["pm"]
    perfect = client.post("/api/quizzes/pm", json={"answers": [q["answer"] for q in qs]}).json()
    assert perfect["pct"] == 100 and perfect["passed_now"] and perfect["review"] == [] and perfect["result"]["passed"]
    wrong = [(q["answer"] + 1) % 4 for q in qs]
    bad = client.post("/api/quizzes/pm", json={"answers": wrong}).json()
    assert bad["pct"] == 0 and not bad["passed_now"] and bad["result"]["best_pct"] == 100 and bad["result"]["passed"] and bad["result"]["attempts"] == 2
    assert all(g["why"] for g in bad["graded"]) and bad["review"]
    assert client.get("/api/quizzes").json()["tracks"]["pm"]["result"]["best_pct"] == 100


@pytest.mark.parametrize("answers", [None, [], [0], "abc", [True] * 6, [9] * 6, [-1] * 6, ["a"] * 6])
def test_bad_quiz_submissions_are_rejected(client, answers):
    assert client.post("/api/quizzes/pm", json={"answers": answers}).status_code == 400


# ------------------------------------------------------------------ search
def test_search_finds_pages_steps_deliverables_and_your_kpis(client):
    client.post("/api/kpis", json={"name": "Fleet margin", "metric": "margin_pct", "direction": "higher", "target": 50, "warn_pct": 10, "window_days": 0})
    kinds = lambda q: {r["kind"] for r in client.get("/api/search", params={"q": q}).json()["results"]}   # noqa: E731
    assert "Manual step" in kinds("okr") and "Deliverable" in kinds("backup") and "Your KPI" in kinds("fleet margin") and "Glossary" in kinds("margin")
    r = client.get("/api/search", params={"q": "restore drill"}).json()["results"]
    assert r and r[0]["href"] == "/tracks/engineering"
    step = next(x for x in client.get("/api/search", params={"q": "okr"}).json()["results"] if x["kind"] == "Manual step")
    assert step["href"].startswith("/tracks/pm?view=manual&step=")


def test_search_empty_and_nonsense(client):
    d = client.get("/api/search").json()
    assert d["suggested"] and d["results"]
    assert client.get("/api/search", params={"q": "zzzzqqqq"}).json()["results"] == []
    assert client.get("/api/search", params={"q": "x" * 101}).status_code == 422


def test_every_search_link_points_at_a_real_page():
    from pathlib import Path
    app = Path(__file__).resolve().parents[2] / "frontend" / "src" / "app"
    for r in search._index():
        top = r["href"].split("?")[0].strip("/")
        assert (app / top / "page.tsx").exists() or (top == "" and (app / "page.tsx").exists()), r["href"]
