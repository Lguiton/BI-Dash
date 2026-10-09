import io

import pytest


# ---------- SCD ----------
def test_scd_type1_vs_type2(client):
    st = client.get("/api/scd/state").json()
    assert len(st["type1"]) == len(st["type2"]) == 4
    base = client.get("/api/scd/compare").json()
    assert all(r["difference"] == 0 for r in base["rows"])  # nothing changed yet

    # ENT-01 (Logistics) has facts on Oct 1, 2, 6. Move it to Fleet effective Oct 4.
    r = client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "Fleet", "effective_date": "2026-10-04"})
    assert r.status_code == 200
    cmp = client.get("/api/scd/compare").json()
    by = {x["category"]: x for x in cmp["rows"]}
    # Type 1 rewrites history: all 3 ENT-01 facts (450+520+540 = 1510) now count as Fleet.
    # Type 2 keeps Oct 1-2 (970) under Logistics and only Oct 6 (540) under Fleet.
    assert by["Logistics"]["type1_revenue"] == 490.0           # only ENT-04 left
    assert by["Logistics"]["type2_revenue"] == 970.0 + 490.0
    assert by["Fleet"]["type1_revenue"] == 610.0 + 1510.0
    assert by["Fleet"]["type2_revenue"] == 610.0 + 540.0
    assert cmp["total_type1"] == cmp["total_type2"]            # same facts, just bucketed differently
    assert cmp["changes_made"] == 1
    t2 = [x for x in client.get("/api/scd/state").json()["type2"] if x["entity_id"] == "ENT-01"]
    assert [x["is_current"] for x in t2] == [False, True]
    assert t2[0]["valid_to"] == "2026-10-03" and t2[1]["valid_from"] == "2026-10-04"


def test_scd_validation_and_reset(client):
    client.get("/api/scd/state")
    assert client.post("/api/scd/change", json={"entity_id": "NOPE", "new_category": "X", "effective_date": "2026-10-04"}).status_code == 404
    assert client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "Logistics", "effective_date": "2026-10-04"}).status_code == 400
    assert client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "Fleet", "effective_date": "1800-01-01"}).status_code == 400
    assert client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "", "effective_date": "2026-10-04"}).status_code == 422
    client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "Fleet", "effective_date": "2026-10-04"})
    # a second change must be after the first one's start date
    assert client.post("/api/scd/change", json={"entity_id": "ENT-01", "new_category": "Express", "effective_date": "2026-10-04"}).status_code == 400
    st = client.post("/api/scd/reset").json()
    assert len(st["type2"]) == 4
    # the real dimension was never touched
    cats = {e["entity_id"]: e["category"] for e in client.get("/api/analytics/meta").json()["entities"]}
    assert cats["ENT-01"] == "Logistics"


def test_scd_tables_are_queryable_in_sql_lab(client):
    client.get("/api/scd/state")
    r = client.post("/api/sql/run", json={"sql": "SELECT COUNT(*) AS n FROM scd_type2 WHERE is_current"})
    assert r.status_code == 200 and r.json()["rows"][0][0] == 4


# ---------- KPIs ----------
def test_kpi_status_logic_and_crud(client):
    assert client.get("/api/kpis").json()["kpis"] == []
    metrics = client.get("/api/kpis/metrics").json()["metrics"]
    assert "margin_pct" in {m["id"] for m in metrics}
    body = {"name": "Revenue goal", "metric": "revenue", "direction": "higher", "target": 3000, "warn_pct": 10, "window_days": 0}
    kid = client.post("/api/kpis", json=body).json()["id"]
    k = client.get("/api/kpis").json()["kpis"][0]
    assert k["value"] == 2990.0 and k["gap"] == -10.0 and k["status"] == "warn"   # within 10% of target
    client.post("/api/kpis", json={**body, "name": "Hit", "target": 2000})
    client.post("/api/kpis", json={**body, "name": "Miss", "target": 5000})
    client.post("/api/kpis", json={**body, "name": "Cost cap", "metric": "cost", "direction": "lower", "target": 1000})
    st = {x["name"]: x["status"] for x in client.get("/api/kpis").json()["kpis"]}
    assert st == {"Revenue goal": "warn", "Hit": "good", "Miss": "bad", "Cost cap": "good"}
    assert client.delete(f"/api/kpis/{kid}").status_code == 200
    assert client.delete(f"/api/kpis/{kid}").status_code == 404


def test_kpi_validation_and_examples(client):
    bad = {"name": "x", "metric": "nope", "direction": "higher", "target": 1}
    assert client.post("/api/kpis", json=bad).status_code == 422
    assert client.post("/api/kpis", json={**bad, "metric": "revenue", "direction": "sideways"}).status_code == 422
    assert client.post("/api/kpis", json={**bad, "metric": "revenue", "name": "  "}).status_code == 422
    assert client.post("/api/kpis", json={**bad, "metric": "revenue", "warn_pct": 500}).status_code == 422
    assert client.post("/api/kpis/examples").json()["added"] == 3
    assert client.post("/api/kpis/examples").status_code == 400  # only for an empty list
    ks = client.get("/api/kpis").json()["kpis"]
    assert len(ks) == 3 and all(k["status"] in {"good", "warn", "bad"} for k in ks)


def test_kpi_margin_is_ratio_of_sums(client):
    client.post("/api/kpis", json={"name": "m", "metric": "margin_pct", "direction": "higher", "target": 70, "window_days": 0})
    k = client.get("/api/kpis").json()["kpis"][0]
    s = client.get("/api/analytics/summary").json()
    assert k["value"] == pytest.approx(s["net_margin_pct"], abs=0.01)


# ---------- Quality ----------
def test_quality_clean_data_and_dirty_data(client):
    q = client.get("/api/quality").json()
    assert q["facts"] == 6 and q["score_pct"] >= 75
    assert {c["id"]: c["status"] for c in q["checks"]}["orphan_facts"] == "ok"
    # add rows with problems straight into the table (the CSV importer would reject most of them)
    from app.services.db import get_cursor
    with get_cursor() as cur:
        cur.execute("INSERT INTO fact_operations VALUES ('Q1','2026-10-01','ENT-01',-5,2,1,10,'Completed')")   # negative + same entity/day
        cur.execute("INSERT INTO fact_operations VALUES ('Q2','2026-10-09','ENT-77',10,20,1,0,'Completed')")   # orphan, loss, zero duration, gap
    q = client.get("/api/quality").json()
    ids = {c["id"]: c for c in q["checks"]}
    assert ids["negative_values"]["status"] == "bad" and ids["negative_values"]["count"] == 1
    assert ids["bad_duration"]["count"] == 1
    assert ids["orphan_facts"]["count"] == 1 and ids["orphan_facts"]["sample"][0]["entity_id"] == "ENT-77"
    assert ids["loss_rows"]["status"] == "warn"
    assert ids["duplicate_business_key"]["count"] == 1
    assert q["score_pct"] < 75


# ---------- Reports ----------
def test_report_xlsx_and_pdf(client):
    x = client.get("/api/report?format=xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "By entity", "Daily", "Insights"]
    assert wb["Summary"]["A5"].value == "Records" and wb["Summary"]["B5"].value == 6
    p = client.get("/api/report?format=pdf&date_from=2026-10-01&date_to=2026-10-06")
    assert p.status_code == 200 and p.content[:5] == b"%PDF-"
    assert client.get("/api/report?format=docx").status_code == 422


def test_report_respects_filters_and_empty(client):
    x = client.get("/api/report?format=xlsx&category=Fleet")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb["Summary"]["B5"].value == 1                     # one Fleet fact (ENT-03, F-104)
    assert "Category: Fleet" in wb["Summary"]["B2"].value
    assert client.get("/api/report?format=pdf&category=Nope").status_code == 400


# ---------- Python tab ----------
def test_notebooks_listing_and_download(client):
    nbs = client.get("/api/python/notebooks").json()["notebooks"]
    assert len(nbs) >= 14 and nbs[0]["file"] == "00_start_here.ipynb"
    assert all(n["title"] and n["concepts"] for n in nbs)
    r = client.get("/api/python/notebooks/05_forecast.ipynb")
    assert r.status_code == 200 and r.json()["nbformat"] == 4
    assert client.get("/api/python/notebooks/nope.ipynb").status_code == 404
    assert client.get("/api/python/notebooks/..%2F..%2Fmain.py").status_code == 404


# ---------- email report ----------
class FakeSMTP:
    sent, logins, started = [], [], []
    def __init__(self, host, port, timeout=None, context=None): self.host, self.port = host, port
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def starttls(self, context=None): FakeSMTP.started.append(True)
    def login(self, u, p): FakeSMTP.logins.append((u, p))
    def send_message(self, m): FakeSMTP.sent.append(m)


def test_email_report_flow(client, monkeypatch):
    import smtplib
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    FakeSMTP.sent.clear(); FakeSMTP.logins.clear(); FakeSMTP.started.clear()
    for k in ("BI_SMTP_HOST", "BI_SMTP_USER", "BI_SMTP_PASSWORD", "BI_REPORT_TO", "BI_SMTP_FROM"):
        monkeypatch.delenv(k, raising=False)
    assert client.get("/api/report/email/status").json()["configured"] is False
    assert client.post("/api/report/email", json={"format": "pdf"}).status_code == 503
    monkeypatch.setenv("BI_SMTP_HOST", "smtp.example.com"); monkeypatch.setenv("BI_SMTP_USER", "me@example.com")
    monkeypatch.setenv("BI_SMTP_PASSWORD", "pw"); monkeypatch.setenv("BI_REPORT_TO", "me@example.com, boss@example.com")
    assert client.get("/api/report/email/status").json()["recipients"] == ["me@example.com", "boss@example.com"]
    r = client.post("/api/report/email", json={"format": "xlsx"})
    assert r.status_code == 200 and r.json()["sent_to"] == ["me@example.com", "boss@example.com"]
    m = FakeSMTP.sent[0]
    assert m["To"] == "me@example.com, boss@example.com" and FakeSMTP.logins == [("me@example.com", "pw")] and FakeSMTP.started
    att = next(m.iter_attachments())
    assert att.get_filename() == "bi_report.xlsx" and att.get_content()[:2] == b"PK"
    assert client.post("/api/report/email", json={"to": "evil@attacker.com"}).status_code == 400    # not on the allow-list
    assert len(FakeSMTP.sent) == 1


def test_email_auth_failure_is_friendly(client, monkeypatch):
    import smtplib
    class Bad(FakeSMTP):
        def login(self, u, p): raise smtplib.SMTPAuthenticationError(535, b"no")
    monkeypatch.setattr(smtplib, "SMTP", Bad)
    for k, v in (("BI_SMTP_HOST", "h"), ("BI_SMTP_USER", "u@x.com"), ("BI_REPORT_TO", "u@x.com")):
        monkeypatch.setenv(k, v)
    r = client.post("/api/report/email", json={})
    assert r.status_code == 502 and "app password" in r.json()["detail"]
