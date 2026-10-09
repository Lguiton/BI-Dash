import io
import json
import xml.etree.ElementTree as ET

import pytest


def get(client, ds, fmt):
    r = client.get(f"/api/export/{ds}", params={"format": fmt})
    assert r.status_code == 200, r.text
    return r


def test_every_format_round_trips_the_same_data(client):
    import csv
    import openpyxl
    import pyarrow.parquet as pq

    expected_ids = [f"F-10{i}" for i in range(1, 7)]
    c = list(csv.DictReader(io.StringIO(get(client, "facts", "csv").text)))
    t = list(csv.DictReader(io.StringIO(get(client, "facts", "tsv").text), delimiter="\t"))
    j = json.loads(get(client, "facts", "json").text)
    x = ET.fromstring(get(client, "facts", "xml").content)
    wb = openpyxl.load_workbook(io.BytesIO(get(client, "facts", "xlsx").content))
    ws = list(wb.active.iter_rows(values_only=True))
    p = pq.read_table(io.BytesIO(get(client, "facts", "parquet").content)).to_pylist()

    assert [r["fact_id"] for r in c] == [r["fact_id"] for r in t] == [r["fact_id"] for r in j] == expected_ids
    assert [r.find("fact_id").text for r in x.findall("row")] == expected_ids
    assert [r[0] for r in ws[1:]] == expected_ids and ws[0][0] == "fact_id"
    assert [r["fact_id"] for r in p] == expected_ids
    # same numbers everywhere
    rev = 450 + 520 + 380 + 610 + 490 + 540
    assert sum(float(r["revenue"]) for r in c) == sum(float(r["revenue"]) for r in t) == rev
    assert sum(r["revenue"] for r in j) == sum(r["revenue"] for r in p) == sum(r[3] for r in ws[1:]) == rev
    assert sum(float(r.find("revenue").text) for r in x.findall("row")) == rev
    # types survive in the typed formats
    import datetime
    assert isinstance(ws[1][1], (datetime.date, datetime.datetime)) and p[0]["record_date"] == datetime.date(2026, 10, 1)


def test_flat_view_and_other_datasets(client):
    import csv
    rows = list(csv.DictReader(io.StringIO(get(client, "operations", "csv").text)))
    assert len(rows) == 6 and {"profit", "margin", "entity_name", "month_name"} <= set(rows[0])
    assert len(list(csv.DictReader(io.StringIO(get(client, "entities", "csv").text)))) == 4
    assert len(list(csv.DictReader(io.StringIO(get(client, "dates", "csv").text)))) == 6


def test_headers_and_errors(client):
    r = get(client, "entities", "xlsx")
    assert r.headers["content-disposition"] == 'attachment; filename="bi_entities.xlsx"'
    assert client.get("/api/export/nope").status_code == 404
    assert client.get("/api/export/facts", params={"format": "pdf"}).status_code == 422
    ds = client.get("/api/export/datasets").json()
    assert {d["id"] for d in ds["datasets"]} == {"operations", "facts", "entities", "dates"} and "parquet" in ds["formats"]
