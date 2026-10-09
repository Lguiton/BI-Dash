"""Download a summary report of the current (filtered) dashboard view as Excel or PDF."""
import io
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from app.routers.analytics import _by_entity, _daily, _kpis, _previous_period, get_insights
from app.services.db import get_cursor
from app.services.filters import Filters, get_filters
from app.services.insights import pct_change

router = APIRouter(prefix="/api/report", tags=["report"])

KPI_ROWS = [
    ("Records", "record_count", "{:,.0f}"), ("Revenue", "total_revenue", "${:,.2f}"),
    ("Operational cost", "total_cost", "${:,.2f}"), ("Net profit", "net_profit", "${:,.2f}"),
    ("Net margin", "net_margin_pct", "{:.2f}%"), ("Units processed", "total_units", "{:,.0f}"),
    ("Revenue per unit", "rev_per_unit", "${:,.2f}"), ("Avg duration (min)", "avg_duration_minutes", "{:,.1f}"),
]


def _describe(flt: Filters) -> str:
    parts = []
    if flt.date_from or flt.date_to:
        parts.append(f"Dates: {flt.date_from or 'start'} to {flt.date_to or 'end'}")
    for label, v in (("Entity", flt.entity_id), ("Category", flt.category), ("Status", flt.status)):
        if v:
            parts.append(f"{label}: {v}")
    return "; ".join(parts) or "All data (no filters)"


def _gather(flt: Filters) -> dict:
    with get_cursor() as cur:
        kpis = _kpis(cur, flt)
        prev_flt = _previous_period(cur, flt)
        prev = _kpis(cur, prev_flt) if prev_flt else None
        entities = _by_entity(cur, flt)
        daily = _daily(cur, flt)
    insights = get_insights(flt)["insights"]
    deltas = {}
    if prev and prev["record_count"]:
        for key in ("total_revenue", "total_cost", "net_profit", "total_units"):
            deltas[key] = pct_change(kpis[key], prev[key])
    return {"kpis": kpis, "prev": prev, "deltas": deltas, "entities": entities, "daily": daily, "insights": insights, "filters": _describe(flt)}


def _fmt(fmt: str, v) -> str:
    return "n/a" if v is None else fmt.format(v)


def _xlsx(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["BI Dashboard report"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append(["View", d["filters"]]); ws.append([])
    ws.append(["KPI", "Value", "vs previous period"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for label, key, fmt in KPI_ROWS:
        ch = d["deltas"].get(key)
        ws.append([label, d["kpis"][key], None if ch is None else f"{ch:+.1f}%"])
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 20

    def sheet(name, cols, rows):
        s = wb.create_sheet(name)
        s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True)
        for r in rows:
            s.append([r.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            s.column_dimensions[s.cell(1, i).column_letter].width = max(14, len(c) + 2)

    sheet("By entity", ["entity_name", "category", "revenue", "profit", "margin_pct", "records", "avg_cost_per_record", "baseline_target", "vs_baseline_pct"],
          [{**e, "avg_cost_per_record": e["avg_cost_per_record"]} for e in d["entities"]])
    sheet("Daily", ["date", "revenue", "cost", "profit", "units"], d["daily"])
    sheet("Insights", ["analytics_type", "severity", "title", "detail"], d["insights"])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _pdf(d: dict) -> bytes:
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    ss = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                            title="BI Dashboard report")
    story = [Paragraph("BI Dashboard report", ss["Title"]), Paragraph(escape(d["filters"]), ss["Normal"]), Spacer(1, 12)]

    def table(rows, widths=None):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                               ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                               ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return t

    k = d["kpis"]
    rows = [["KPI", "Value", "vs previous period"]]
    for label, key, fmt in KPI_ROWS:
        ch = d["deltas"].get(key)
        rows.append([label, _fmt(fmt, k[key]), "" if ch is None else f"{ch:+.1f}%"])
    story += [table(rows, [2.4 * inch, 1.6 * inch, 1.6 * inch]), Spacer(1, 14)]

    if len(d["daily"]) >= 2:
        story.append(Paragraph("Daily revenue", ss["Heading3"]))
        pts = [(i, r["revenue"]) for i, r in enumerate(d["daily"])]
        dr = Drawing(480, 150)
        lp = LinePlot(); lp.x, lp.y, lp.width, lp.height = 40, 20, 420, 120
        lp.data = [pts]; lp.lines[0].strokeColor = colors.HexColor("#2563eb"); lp.lines[0].strokeWidth = 1.2
        lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = 0, len(pts) - 1
        lp.xValueAxis.labelTextFormat = lambda v, ds=d["daily"]: ds[int(v)]["date"][5:] if 0 <= int(v) < len(ds) else ""
        lp.xValueAxis.labels.fontSize = 7; lp.yValueAxis.labels.fontSize = 7
        dr.add(lp); story += [dr, Spacer(1, 10)]

    if d["entities"]:
        story.append(Paragraph("By entity (top 15 by revenue)", ss["Heading3"]))
        rows = [["Entity", "Category", "Revenue", "Margin", "Cost vs budget"]]
        for e in d["entities"][:15]:
            rows.append([e["entity_name"][:28], e["category"], f"${e['revenue']:,.0f}", f"{e['margin_pct']:.1f}%",
                         "n/a" if e["vs_baseline_pct"] is None else f"{e['vs_baseline_pct']:+.1f}%"])
        story += [table(rows, [2.0 * inch, 1.2 * inch, 1.1 * inch, 0.9 * inch, 1.2 * inch]), Spacer(1, 14)]

    story.append(Paragraph("Insights", ss["Heading3"]))
    for ins in d["insights"][:14]:
        story.append(Paragraph(f"<b>[{escape(ins['analytics_type'])}]</b> {escape(ins['title'])}. {escape(ins.get('detail', ''))}", ss["BodyText"]))
    doc.build(story)
    return buf.getvalue()


@router.get("")
def download(format: Literal["xlsx", "pdf"] = Query("xlsx"), flt: Filters = Depends(get_filters)):
    d = _gather(flt)
    if not d["kpis"]["record_count"]:
        raise HTTPException(400, "No data matches the current filters, so there is nothing to report.")
    try:
        data = _xlsx(d) if format == "xlsx" else _pdf(d)
    except ImportError as e:
        raise HTTPException(501, f"Report export needs an extra package: pip install openpyxl reportlab ({e.name})")
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if format == "xlsx" else "application/pdf"
    return Response(data, media_type=mime, headers={"Content-Disposition": f'attachment; filename="bi_report.{format}"'})


# ---------------- email the report ----------------
# Safe-by-design: mail only goes to addresses YOU list in BI_REPORT_TO, so the endpoint can't be used as an open mail relay.
import os
import smtplib
import ssl
from email.message import EmailMessage

from pydantic import BaseModel


def _mail_config() -> dict:
    from app.services import mailer
    return mailer.config()


class EmailIn(BaseModel):
    format: Literal["xlsx", "pdf"] = "pdf"
    to: str | None = None


@router.get("/email/status")
def email_status():
    c = _mail_config()
    ready = bool(c["host"] and c["to"] and c["from"])
    return {"configured": ready, "recipients": c["to"],
            "hint": "" if ready else "Set BI_SMTP_HOST, BI_SMTP_USER, BI_SMTP_PASSWORD and BI_REPORT_TO (the addresses allowed to receive reports) in backend/.env, then restart."}


@router.post("/email")
def email_report(body: EmailIn, flt: Filters = Depends(get_filters)):
    c = _mail_config()
    if not (c["host"] and c["to"] and c["from"]):
        raise HTTPException(503, "Email isn't configured. Set BI_SMTP_HOST, BI_SMTP_USER, BI_SMTP_PASSWORD and BI_REPORT_TO in backend/.env.")
    recipients = c["to"]
    if body.to:
        if body.to.lower() not in [a.lower() for a in c["to"]]:
            raise HTTPException(400, "That address isn't in BI_REPORT_TO, so reports can't be sent to it.")
        recipients = [body.to]
    d = _gather(flt)
    if not d["kpis"]["record_count"]:
        raise HTTPException(400, "No data matches the current filters, so there is nothing to report.")
    data = _xlsx(d) if body.format == "xlsx" else _pdf(d)
    k = d["kpis"]
    msg = EmailMessage()
    msg["Subject"] = f"BI report: revenue ${k['total_revenue']:,.0f}, margin {k['net_margin_pct']:.1f}%"
    msg["From"], msg["To"] = c["from"], ", ".join(recipients)
    msg.set_content(f"Your BI report is attached.\n\nScope: {d['filters']}\nRecords: {k['record_count']:,}\n"
                    f"Revenue: ${k['total_revenue']:,.2f}\nNet profit: ${k['net_profit']:,.2f}\nNet margin: {k['net_margin_pct']:.2f}%\n")
    maintype, subtype = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet") if body.format == "xlsx" else ("application", "pdf")
    msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=f"bi_report.{body.format}")
    try:
        if c["security"] == "ssl":
            server = smtplib.SMTP_SSL(c["host"], c["port"], timeout=20, context=ssl.create_default_context())
        else:
            server = smtplib.SMTP(c["host"], c["port"], timeout=20)
        with server:
            if c["security"] == "starttls":
                server.starttls(context=ssl.create_default_context())
            if c["user"]:
                server.login(c["user"], c["password"])
            server.send_message(msg)
    except smtplib.SMTPAuthenticationError:
        raise HTTPException(502, "The mail server rejected the username or password (Gmail needs an app password).")
    except (smtplib.SMTPException, OSError) as e:
        raise HTTPException(502, f"Couldn't send the email ({type(e).__name__}). Check BI_SMTP_HOST, port and security.")
    from app.services import state
    state.audit("email_report", f"{body.format} sent to {len(recipients)} recipient(s)")
    return {"sent_to": recipients, "format": body.format}
