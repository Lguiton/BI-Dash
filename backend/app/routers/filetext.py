"""Turn an Excel, PDF or Word file into plain text so the Upload button and the agents can read it. Nothing is stored."""
from __future__ import annotations

import csv
import io
import re
import zipfile

from fastapi import APIRouter, File, HTTPException, UploadFile

from app.services import fileformats

router = APIRouter(prefix="/api/files", tags=["files"])
MAX = 25 * 1024 * 1024


def _xlsx(raw: bytes) -> tuple[str, str]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Couldn't open that Excel file. Is it damaged or password protected?") from None
    out, note = [], ""
    for i, ws in enumerate(wb.worksheets):
        if i >= 6:
            note = "Only the first 6 sheets were read. "
            break
        buf = io.StringIO()
        w = csv.writer(buf)
        n = 0
        for row in ws.iter_rows(values_only=True):
            if n >= 2000:
                note += f"Sheet '{ws.title}' was cut at 2,000 rows. "
                break
            if any(c is not None for c in row):
                w.writerow(["" if c is None else (round(c, 4) if isinstance(c, float) else c) for c in row])
                n += 1
        out.append(f"## Sheet: {ws.title}\n{buf.getvalue()}")
    return "\n".join(out), note


def _docx(raw: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except Exception:
        raise HTTPException(400, "Couldn't open that Word file.") from None
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")


@router.post("/text")
async def to_text(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    raw = await file.read(MAX + 1)
    if len(raw) > MAX:
        raise HTTPException(400, "That file is over 25 MB. Export a smaller slice.")
    note = ""
    if name.endswith((".xlsx", ".xlsm")):
        text, note = _xlsx(raw)
    elif name.endswith(".pdf"):
        text = fileformats.pdf_text(raw)
    elif name.endswith(".docx"):
        text = _docx(raw)
    else:
        raise HTTPException(400, "This reader handles .xlsx, .xlsm, .pdf and .docx files.")
    if not text.strip():
        raise HTTPException(400, "No readable text found (a scanned PDF needs OCR first).")
    return {"name": file.filename, "chars": len(text), "text": text[:60000], "note": (note + ("Text was cut at 60,000 characters." if len(text) > 60000 else "")).strip()}
