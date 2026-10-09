"""Export the engagement plan and status report as Excel or PDF, so the work can be shown to someone else."""
from __future__ import annotations

import io
from xml.sax.saxutils import escape

from app.services import company

DISC = {d[0]: d[1] for d in company.DISCIPLINES}


def _status(i: dict) -> str:
    if i["done"]:
        return "Done (detected from data)" if i["detected"] else "Done (ticked by you)"
    return "OVERDUE" if i["overdue"] else "Open"


def _data() -> dict:
    ov = company.overview()
    return {"ov": ov, "snaps": company.snapshots(12), "items": [i for p in ov["phases"] for i in p["deliverables"]]}


def xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    d = _data()
    ov = d["ov"]
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s.append(["Company engagement: status report"]); s["A1"].font = Font(bold=True, size=14)
    s.append(["Workspace", ov["workspace"]])
    b = ov["brief"]
    for k, label in (("company", "Company"), ("goal", "Goal"), ("notes", "Notes")):
        s.append([label, b.get(k, "")])
    s.append(["Done", f"{ov['done']} of {ov['total']} ({ov['pct']}%)"])
    s.append(["Overdue", ov["overdue"]]); s.append(["Due within 7 days", ov["due_soon"]])
    if d["snaps"]:
        s.append([]); s.append(["Latest brief", d["snaps"][0]["brief"]])
    s.column_dimensions["A"].width = 22; s.column_dimensions["B"].width = 90

    def sheet(name, cols, rows, widths):
        w = wb.create_sheet(name)
        w.append(cols)
        for c in w[1]:
            c.font = Font(bold=True)
        for r in rows:
            w.append(r)
        for i, wd in enumerate(widths, 1):
            w.column_dimensions[w.cell(1, i).column_letter].width = wd
        return w

    ph = {p["id"]: p["name"] for p in ov["phases"]}
    pid = {i["id"]: p["id"] for p in ov["phases"] for i in p["deliverables"]}
    w = sheet("Deliverables", ["Phase", "Discipline", "Deliverable", "Status", "Evidence", "Due", "Note", "Why it matters"],
              [[ph[pid[i["id"]]], DISC[i["discipline"]], i["title"], _status(i), i["detail"], i["due"], i["note"], i["why"]] for i in d["items"]],
              [11, 30, 52, 26, 40, 12, 40, 50])
    red = PatternFill("solid", fgColor="FDE2E2")
    for row in w.iter_rows(min_row=2):
        if row[3].value == "OVERDUE":
            for c in row:
                c.fill = red
    sheet("Disciplines", ["Discipline", "Deliverables done", "Deliverables total", "Learning steps done", "Learning steps total"],
          [[x["name"], x["done"], x["total"], x["learning_done"], x["learning_total"]] for x in ov["disciplines"]], [38, 18, 18, 20, 20])
    sheet("Snapshots", ["When (UTC)", "Kind", "Done", "Total", "Percent", "Brief"],
          [[x["at"], x["kind"], x["done"], x["total"], x["pct"], x["brief"]] for x in d["snaps"]], [20, 10, 8, 8, 9, 110])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def pdf() -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    d = _data()
    ov = d["ov"]
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=8, leading=10)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title="Company engagement: status report")
    story = [Paragraph("Company engagement: status report", ss["Title"])]
    b = ov["brief"]
    for k, label in (("company", "Company"), ("goal", "Goal")):
        if b.get(k):
            story.append(Paragraph(f"<b>{label}:</b> {escape(b[k])}", ss["BodyText"]))
    story.append(Paragraph(f"<b>Progress:</b> {ov['done']} of {ov['total']} deliverables done ({ov['pct']}%). Overdue: {ov['overdue']}. Due within 7 days: {ov['due_soon']}.", ss["BodyText"]))
    if d["snaps"]:
        story += [Spacer(1, 6), Paragraph("<b>Latest brief:</b> " + escape(d["snaps"][0]["brief"]), ss["BodyText"])]

    def table(rows, widths):
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                               ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story += [Spacer(1, 10), Paragraph("By discipline", ss["Heading3"]),
              table([["Discipline", "Deliverables", "Learning steps"]] + [[x["name"], f"{x['done']}/{x['total']}", f"{x['learning_done']}/{x['learning_total']}"] for x in ov["disciplines"]],
                    [3.4 * inch, 1.6 * inch, 1.6 * inch])]
    for p in ov["phases"]:
        story += [Spacer(1, 8), Paragraph(f"{escape(p['name'])}: {p['done']}/{p['total']}", ss["Heading3"])]
        rows = [["Deliverable", "Discipline", "Status", "Due", "Note"]]
        for i in p["deliverables"]:
            rows.append([Paragraph(escape(i["title"]), small), Paragraph(escape(DISC[i["discipline"]]), small), Paragraph(_status(i), small), i["due"] or "", Paragraph(escape(i["note"]), small)])
        story.append(table(rows, [2.5 * inch, 1.4 * inch, 1.2 * inch, 0.8 * inch, 1.4 * inch]))
    story += [Spacer(1, 10), Paragraph("Deliverables marked 'detected' are proven by the workspace's own data. The rest are ticked by hand. This report is a self-assessment, not an audit.", small)]
    doc.build(story)
    return buf.getvalue()
