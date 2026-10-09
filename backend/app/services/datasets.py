"""Generic tables: import any CSV/Excel file as its own table, profile it, and run safe chart queries on it.

Everything the user supplies (table names, column names) is either sanitized or checked against the table's real
columns and then double-quoted, so it can never change the shape of a query. Imports are staged and swapped in a
transaction, so a failed import leaves the previous table untouched.
"""
import csv
import io
import json
import math
import re
from datetime import date, datetime

import pyarrow as pa

from app.config import MAX_TABLE_ROWS
from app.services.db import fetch_all, fetch_one, get_cursor

NULL_TOKENS = {"", "na", "n/a", "null", "none", "nan", "#n/a", "-"}
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d-%b-%Y", "%b %d, %Y")
TS_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M")
INT_RE = re.compile(r"^[+-]?(0|[1-9]\d{0,17})$")
NUM_RE = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
MONEY_RE = re.compile(r"^[+-]?\$?\s?\d{1,3}(,\d{3})+(\.\d+)?$|^[+-]?\$\s?\d+(\.\d+)?$")
RESERVED = {
    "dim_entities", "fact_operations", "app_meta", "dim_date", "v_operations_flat", "user_tables", "dim_entities_scd",
    "kpi_definitions", "study_progress", "ml_runs", "ai_log",
}
MAX_COLUMNS = 300
AGGS = {"sum": "SUM", "avg": "AVG", "min": "MIN", "max": "MAX", "count": "COUNT", "median": "MEDIAN", "distinct": "COUNT(DISTINCT"}
BUCKETS = ("day", "week", "month", "quarter", "year")


class DataError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message = message
        self.status = status


def q(ident: str) -> str:
    return '"' + ident.replace('"', '""') + '"'


# ---------------------------------------------------------------- registry
def ensure_registry() -> None:
    with get_cursor() as cur:
        cur.execute(
            """CREATE TABLE IF NOT EXISTS user_tables (
                table_name VARCHAR PRIMARY KEY,
                label VARCHAR,
                source VARCHAR,
                rows_count BIGINT,
                columns_json VARCHAR,
                created_at TIMESTAMP,
                updated_at TIMESTAMP
            )"""
        )


def list_tables() -> list[dict]:
    ensure_registry()
    with get_cursor() as cur:
        rows = fetch_all(cur, "SELECT table_name, label, source, rows_count, columns_json, CAST(created_at AS VARCHAR) AS created_at, CAST(updated_at AS VARCHAR) AS updated_at FROM user_tables ORDER BY updated_at DESC")
    for r in rows:
        r["columns"] = json.loads(r.pop("columns_json") or "[]")
    return rows


def table_columns(table: str) -> list[dict]:
    """Real columns of a registered table; raises if the table is not one of ours."""
    ensure_registry()
    with get_cursor() as cur:
        if not cur.execute("SELECT 1 FROM user_tables WHERE table_name = ?", [table]).fetchone():
            raise DataError(f"No imported table named '{table}'.", 404)
        return fetch_all(cur, "SELECT column_name AS name, data_type AS type FROM information_schema.columns WHERE table_schema = 'main' AND table_name = ? ORDER BY ordinal_position", [table])


def kind_of(duck_type: str) -> str:
    t = duck_type.upper()
    if t in ("BOOLEAN",):
        return "bool"
    if t == "DATE" or t.startswith("TIMESTAMP"):
        return "date"
    if t in ("BIGINT", "INTEGER", "SMALLINT", "TINYINT", "HUGEINT", "UBIGINT", "UINTEGER") or t.startswith("DECIMAL") or t in ("DOUBLE", "FLOAT", "REAL"):
        return "number"
    return "text"


def delete_table(table: str) -> None:
    table_columns(table)
    with get_cursor() as cur:
        cur.execute("BEGIN")
        try:
            cur.execute(f"DROP TABLE IF EXISTS {q(table)}")
            cur.execute("DELETE FROM user_tables WHERE table_name = ?", [table])
            cur.execute("COMMIT")
        except Exception:
            cur.execute("ROLLBACK")
            raise


# ---------------------------------------------------------------- reading files
def decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeError:
            continue
    return raw.decode("latin-1")


def sniff_delimiter(text: str) -> str:
    head = "\n".join(text.splitlines()[:20])
    best, best_n = ",", 0
    for d in (",", ";", "\t", "|"):
        counts = [len(r) for r in csv.reader(io.StringIO(head), delimiter=d)]
        if counts and min(counts) > 1 and len(set(counts)) == 1 and counts[0] > best_n:
            best, best_n = d, counts[0]
    if best_n == 0:  # ragged file: pick the delimiter that appears most in the header line
        first = head.splitlines()[0] if head else ""
        best = max((",", ";", "\t", "|"), key=first.count)
    return best


def read_table(raw: bytes, filename: str = "data.csv", sheet: str | None = None) -> tuple[list[str], list[list[str | None]], dict]:
    """Return (header names as written, rows of strings or None, info). Raises DataError on unreadable files."""
    if not raw.strip():
        raise DataError("The file is empty.")
    info: dict = {"filename": filename}
    if filename.lower().endswith((".xlsx", ".xlsm")):
        rows, info = _read_xlsx(raw, sheet, info)
    elif filename.lower().endswith(".xls"):
        raise DataError("Old .xls files are not supported. Save the sheet as .xlsx or CSV and try again.")
    elif filename.lower().endswith(".json"):
        rows = _read_json(raw)
    else:
        text = decode(raw)
        delim = sniff_delimiter(text)
        info["delimiter"] = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}[delim]
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise DataError("Need a header row and at least one data row.")
    header, body = rows[0], rows[1:]
    if len(body) > MAX_TABLE_ROWS:
        raise DataError(f"{len(body):,} rows is over the {MAX_TABLE_ROWS:,}-row limit.")
    width = len(header)
    if width > MAX_COLUMNS:
        raise DataError(f"{width} columns is over the {MAX_COLUMNS}-column limit.")
    out = []
    for r in body:
        r = [None if (c is None or c.strip().lower() in NULL_TOKENS) else c.strip() for c in r[:width]]
        r += [None] * (width - len(r))
        out.append(r)
    return header, out, info


def _read_xlsx(raw: bytes, sheet: str | None, info: dict):
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise DataError("Excel support needs openpyxl (pip install openpyxl).") from e
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise DataError("Could not read that Excel file. Is it a valid .xlsx?") from e
    info["sheets"] = wb.sheetnames
    name = sheet if sheet in wb.sheetnames else wb.sheetnames[0]
    info["sheet"] = name
    rows = []
    for row in wb[name].iter_rows(values_only=True):
        rows.append([_cell(v) for v in row])
        if len(rows) > MAX_TABLE_ROWS + 1:
            break
    wb.close()
    return rows, info


def _cell(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat(sep=" ") if (v.hour or v.minute or v.second) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _read_json(raw: bytes) -> list[list[str | None]]:
    try:
        data = json.loads(decode(raw))
    except ValueError as e:
        raise DataError("That file is not valid JSON.") from e
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        data = lists[0] if lists else [data]
    if not isinstance(data, list) or not data or not all(isinstance(x, dict) for x in data):
        raise DataError("JSON must be a list of objects, e.g. [{\"a\": 1}, {\"a\": 2}].")
    cols: list[str] = []
    for obj in data[:5000]:
        for k in obj:
            if k not in cols:
                cols.append(k)
    return [cols] + [[None if obj.get(c) is None else (json.dumps(obj[c]) if isinstance(obj.get(c), (dict, list)) else _cell(obj.get(c))) for c in cols] for obj in data]


# ---------------------------------------------------------------- names and types
def slug(name: str, fallback: str = "column") -> str:
    s = re.sub(r"[^a-z0-9]+", "_", (name or "").strip().lower()).strip("_")
    if not s:
        s = fallback
    if s[0].isdigit():
        s = "c_" + s
    return s[:60]


def sanitize_columns(header: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(header):
        s = slug(h, f"column_{i + 1}")
        n = seen.get(s, 0)
        seen[s] = n + 1
        out.append(s if n == 0 else f"{s}_{n + 1}")
    return out


def table_name_for(label: str) -> str:
    name = slug(label, "table")
    if name in RESERVED or name.startswith(("v_", "dim_", "fact_", "sqlite_", "duckdb_")):
        name = "my_" + name
    return name


def _parse_date(s: str, fmt: str) -> date:
    return datetime.strptime(s, fmt).date()


def infer_column(values: list[str | None]) -> tuple[str, str | None]:
    """Return (type, date format). Type is one of bool, int, float, date, timestamp, varchar."""
    vals = [v for v in values if v is not None]
    if not vals:
        return "varchar", None
    low = {v.lower() for v in vals}
    if low <= {"true", "false"}:
        return "bool", None
    if all(INT_RE.match(v) for v in vals):
        return "int", None
    def is_num(v: str) -> bool:
        digits = v.lstrip("+-")
        if len(digits) > 1 and digits[0] == "0" and digits[1].isdigit():   # 007, 00123: codes, keep as text
            return False
        return bool(NUM_RE.match(v) or MONEY_RE.match(v))

    if all(is_num(v) for v in vals):
        try:
            if all(math.isfinite(float(v.replace("$", "").replace(",", "").replace(" ", ""))) for v in vals):
                return "float", None
        except ValueError:
            pass
    sample = vals
    for fmt in DATE_FORMATS:
        try:
            for v in sample:
                datetime.strptime(v, fmt)
            return "date", fmt
        except ValueError:
            continue
    for fmt in TS_FORMATS:
        try:
            for v in sample:
                datetime.strptime(v, fmt)
            return "timestamp", fmt
        except ValueError:
            continue
    try:  # ISO timestamps with fractions / timezone
        for v in sample:
            datetime.fromisoformat(v.replace("Z", "+00:00"))
        return "timestamp", "iso"
    except ValueError:
        pass
    return "varchar", None


def infer_types(columns: list[str], rows: list[list[str | None]]) -> list[dict]:
    out = []
    for i, name in enumerate(columns):
        t, fmt = infer_column([r[i] for r in rows])
        out.append({"name": name, "type": t, "fmt": fmt})
    return out


def _convert(col: dict, v: str | None):
    if v is None:
        return None
    t = col["type"]
    if t == "bool":
        return v.lower() == "true"
    if t == "int":
        return int(v)
    if t == "float":
        return float(v.replace("$", "").replace(",", "").replace(" ", ""))
    if t == "date":
        return _parse_date(v, col["fmt"])
    if t == "timestamp":
        if col["fmt"] == "iso":
            d = datetime.fromisoformat(v.replace("Z", "+00:00"))
            return d.replace(tzinfo=None)
        return datetime.strptime(v, col["fmt"])
    return v


ARROW = {"bool": pa.bool_(), "int": pa.int64(), "float": pa.float64(), "date": pa.date32(), "timestamp": pa.timestamp("us"), "varchar": pa.string()}


def preview(raw: bytes, filename: str, sheet: str | None = None, limit: int = 15) -> dict:
    header, rows, info = read_table(raw, filename, sheet)
    names = sanitize_columns(header)
    cols = infer_types(names, rows)
    return {
        **info,
        "rows_total": len(rows),
        "header": header,
        "columns": [{"name": c["name"], "original": header[i], "type": c["type"], "samples": [r[i] for r in rows[:3] if r[i] is not None][:3]} for i, c in enumerate(cols)],
        "preview": [dict(zip(names, r)) for r in rows[:limit]],
        "suggested_name": table_name_for(re.sub(r"\.[A-Za-z0-9]+$", "", filename)),
    }


# ---------------------------------------------------------------- import
def _arrow_table(cols: list[dict], rows: list[list[str | None]]) -> pa.Table:
    arrays = []
    for i, c in enumerate(cols):
        try:
            data = [_convert(c, r[i]) for r in rows]
        except (ValueError, OverflowError) as e:  # pragma: no cover - inference makes this unreachable
            raise DataError(f"Column '{c['name']}' could not be converted: {e}") from e
        arrays.append(pa.array(data, type=ARROW[c["type"]]))
    return pa.Table.from_arrays(arrays, names=[c["name"] for c in cols])


def import_rows(header: list[str], rows: list[list[str | None]], table: str, mode: str = "replace", source: str = "upload",
                min_ratio: float | None = None) -> dict:
    """Create/replace/append a table. `min_ratio` refuses a replace that would shrink the table below that share of rows."""
    ensure_registry()
    if mode not in ("create", "replace", "append"):
        raise DataError("mode must be create, replace or append.")
    names = sanitize_columns(header)
    cols = infer_types(names, rows)
    arrow = _arrow_table(cols, rows)
    label = table
    table = table_name_for(table)
    with get_cursor() as cur:
        registered = cur.execute("SELECT rows_count FROM user_tables WHERE table_name = ?", [table]).fetchone()
        created = cur.execute("SELECT created_at FROM user_tables WHERE table_name = ?", [table]).fetchone()
        exists = cur.execute("SELECT 1 FROM information_schema.tables WHERE table_schema='main' AND table_name = ?", [table]).fetchone()
        if exists and not registered:
            raise DataError(f"'{table}' is already a built-in table. Pick another name.")
        if registered and mode == "create":
            raise DataError(f"A table named '{table}' already exists. Choose replace or append, or a new name.")
        if mode == "replace" and registered and min_ratio and registered[0] and len(rows) < registered[0] * min_ratio:
            raise DataError(f"Refusing to replace '{table}': the new data has {len(rows):,} rows but the table has {registered[0]:,}. "
                            f"That is under {int(min_ratio * 100)}% of the old size, which usually means a broken export.", 409)
        cur.register("_incoming", arrow)
        cur.execute("BEGIN")
        try:
            if mode == "append" and registered:
                have = [r[0] for r in cur.execute("SELECT column_name FROM information_schema.columns WHERE table_schema='main' AND table_name = ? ORDER BY ordinal_position", [table]).fetchall()]
                if have != names:
                    raise DataError(f"Columns don't match the existing table. Existing: {', '.join(have)}. New file: {', '.join(names)}.")
                cur.execute(f"INSERT INTO {q(table)} SELECT * FROM _incoming")
            else:
                cur.execute(f"CREATE OR REPLACE TABLE {q(table)} AS SELECT * FROM _incoming")
            total = cur.execute(f"SELECT COUNT(*) FROM {q(table)}").fetchone()[0]
            col_meta = fetch_all(cur, "SELECT column_name AS name, data_type AS type FROM information_schema.columns WHERE table_schema='main' AND table_name = ? ORDER BY ordinal_position", [table])
            cur.execute("DELETE FROM user_tables WHERE table_name = ?", [table])
            cur.execute("INSERT INTO user_tables VALUES (?, ?, ?, ?, ?, COALESCE(?, now()), now())",
                        [table, label, source, total, json.dumps(col_meta), created[0] if created else None])
            cur.execute("COMMIT")
        except Exception:
            cur.execute("ROLLBACK")
            raise
        finally:
            try:
                cur.unregister("_incoming")
            except Exception:  # noqa: BLE001
                pass
    return {"table": table, "rows": total, "added": len(rows), "mode": mode, "columns": col_meta}


def import_file(raw: bytes, filename: str, table: str | None, mode: str = "replace", sheet: str | None = None, source: str = "upload") -> dict:
    header, rows, _ = read_table(raw, filename, sheet)
    name = table or re.sub(r"\.[A-Za-z0-9]+$", "", filename)
    return import_rows(header, rows, name, mode, source)


# ---------------------------------------------------------------- profile and rows
def profile(table: str) -> dict:
    cols = table_columns(table)
    with get_cursor() as cur:
        total = cur.execute(f"SELECT COUNT(*) FROM {q(table)}").fetchone()[0]
        out = []
        for c in cols:
            k = kind_of(c["type"])
            col = q(c["name"])
            base = fetch_one(cur, f"SELECT COUNT(*) - COUNT({col}) AS nulls, COUNT(DISTINCT {col}) AS distinct_count FROM {q(table)}")
            item = {"name": c["name"], "type": c["type"], "kind": k, **base}
            if k == "number":
                item.update(fetch_one(cur, f"SELECT MIN({col}) AS min, MAX({col}) AS max, AVG({col}) AS mean, MEDIAN({col}) AS median FROM {q(table)}"))
            elif k == "date":
                item.update({kk: str(vv) if vv is not None else None for kk, vv in fetch_one(cur, f"SELECT MIN({col}) AS min, MAX({col}) AS max FROM {q(table)}").items()})
            elif k == "text":
                item["top"] = fetch_all(cur, f"SELECT CAST({col} AS VARCHAR) AS value, COUNT(*) AS n FROM {q(table)} WHERE {col} IS NOT NULL GROUP BY 1 ORDER BY n DESC, 1 LIMIT 5")
            out.append(item)
    return {"table": table, "rows": total, "columns": out}


def sample_rows(table: str, limit: int = 20, offset: int = 0) -> dict:
    table_columns(table)
    limit = max(1, min(limit, 200))
    with get_cursor() as cur:
        rows = fetch_all(cur, f"SELECT * FROM {q(table)} LIMIT {int(limit)} OFFSET {max(0, int(offset))}")
    return {"table": table, "rows": [{k: (str(v) if isinstance(v, (date, datetime)) else v) for k, v in r.items()} for r in rows]}


# ---------------------------------------------------------------- chart queries
def _col(table_cols: list[dict], name: str | None, kinds: tuple[str, ...] | None = None, what: str = "column") -> dict:
    for c in table_cols:
        if c["name"] == name:
            if kinds and kind_of(c["type"]) not in kinds:
                raise DataError(f"'{name}' is a {kind_of(c['type'])} column; this chart needs {' or '.join(kinds)}.")
            return c
    raise DataError(f"Unknown {what} '{name}'.")


def _json(v):
    if isinstance(v, (date, datetime)):
        return str(v)[:10] if isinstance(v, date) and not isinstance(v, datetime) else str(v)
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if hasattr(v, "quantize"):  # Decimal
        return float(v)
    return v


def aggregate(table: str, x: str, y: str | None, agg: str = "sum", bucket: str | None = None, group: str | None = None, limit: int = 30,
              sort: str = "auto") -> dict:
    """Group by x (optionally a date bucket) and aggregate y. `group` splits into series (max 8 shown)."""
    cols = table_columns(table)
    xc = _col(cols, x, what="x column")
    agg = agg.lower()
    if agg not in AGGS:
        raise DataError(f"agg must be one of {', '.join(AGGS)}.")
    if agg == "count" or y in (None, "", "*"):
        measure, label = "COUNT(*)", "count"
    else:
        yc = _col(cols, y, ("number", "bool") if agg not in ("count", "distinct", "min", "max") else None, "y column")
        yq = q(yc["name"])
        measure = "COUNT(DISTINCT " + yq + ")" if agg == "distinct" else f"{AGGS[agg]}({yq})"
        label = f"{agg}({y})"
    xk = kind_of(xc["type"])
    xq = q(xc["name"])
    if xk == "date" and bucket:
        if bucket not in BUCKETS:
            raise DataError(f"bucket must be one of {', '.join(BUCKETS)}.")
        xexpr = f"CAST(date_trunc('{bucket}', {xq}) AS DATE)"
    else:
        xexpr = f"CAST({xq} AS VARCHAR)" if xk == "text" else xq
    limit = max(1, min(limit, 500))
    order = "1" if (xk == "date" or sort == "x") else "2 DESC"
    with get_cursor() as cur:
        if group:
            gc = _col(cols, group, what="group column")
            gq = q(gc["name"])
            top = [r[0] for r in cur.execute(f"SELECT CAST({gq} AS VARCHAR) FROM {q(table)} WHERE {gq} IS NOT NULL GROUP BY 1 ORDER BY COUNT(*) DESC LIMIT 8").fetchall()]
            rows = fetch_all(cur, f"SELECT {xexpr} AS x, CAST({gq} AS VARCHAR) AS g, {measure} AS v FROM {q(table)} WHERE {xq} IS NOT NULL AND CAST({gq} AS VARCHAR) IN ({','.join('?' * len(top)) or 'NULL'}) GROUP BY 1, 2 ORDER BY 1, 2", top)
            xs = []
            for r in rows:
                if r["x"] not in xs:
                    xs.append(r["x"])
            xs = xs if xk == "date" else xs[:limit]
            by = {(str(_json(r["x"])), r["g"]): _json(r["v"]) for r in rows}
            data = [{"x": str(_json(x_)), **{g: by.get((str(_json(x_)), g)) for g in top}} for x_ in xs[-limit:] if True]
            return {"x": x, "measure": label, "series": top, "data": data, "truncated_series": False}
        rows = fetch_all(cur, f"SELECT {xexpr} AS x, {measure} AS v FROM {q(table)} WHERE {xq} IS NOT NULL GROUP BY 1 ORDER BY {order} LIMIT {limit}")
        if xk == "date":
            rows = rows[-limit:]
    return {"x": x, "measure": label, "series": ["v"], "data": [{"x": str(_json(r["x"])), "v": _json(r["v"])} for r in rows]}


def histogram(table: str, column: str, bins: int = 20) -> dict:
    cols = table_columns(table)
    c = _col(cols, column, ("number",))
    bins = max(2, min(bins, 100))
    cq = q(c["name"])
    with get_cursor() as cur:
        mm = fetch_one(cur, f"SELECT MIN({cq}) AS lo, MAX({cq}) AS hi, COUNT({cq}) AS n FROM {q(table)}")
        if not mm["n"]:
            return {"column": column, "bins": [], "n": 0}
        lo, hi = float(mm["lo"]), float(mm["hi"])
        if lo == hi:
            return {"column": column, "n": mm["n"], "bins": [{"start": lo, "end": hi, "label": f"{lo:g}", "count": mm["n"]}]}
        width = (hi - lo) / bins
        rows = fetch_all(cur, f"SELECT LEAST(CAST(FLOOR((CAST({cq} AS DOUBLE) - ?) / ?) AS INTEGER), ?) AS b, COUNT(*) AS n FROM {q(table)} WHERE {cq} IS NOT NULL GROUP BY 1 ORDER BY 1", [lo, width, bins - 1])
    counts = {r["b"]: r["n"] for r in rows}
    out = []
    for i in range(bins):
        s, e = lo + i * width, lo + (i + 1) * width
        out.append({"start": s, "end": e, "label": f"{s:.4g}", "count": counts.get(i, 0)})
    return {"column": column, "n": mm["n"], "bins": out}


def box(table: str, column: str, by: str | None = None) -> dict:
    cols = table_columns(table)
    c = _col(cols, column, ("number",))
    cq = q(c["name"])
    stat = (f"MIN({cq}) AS mn, quantile_cont({cq}, 0.25) AS q1, MEDIAN({cq}) AS med, quantile_cont({cq}, 0.75) AS q3, MAX({cq}) AS mx, "
            f"AVG({cq}) AS mean, COUNT({cq}) AS n")
    with get_cursor() as cur:
        if by:
            g = _col(cols, by, what="group column")
            gq = q(g["name"])
            rows = fetch_all(cur, f"SELECT CAST({gq} AS VARCHAR) AS name, {stat} FROM {q(table)} WHERE {cq} IS NOT NULL AND {gq} IS NOT NULL GROUP BY 1 ORDER BY n DESC, 1 LIMIT 12")
        else:
            rows = [{"name": column, **fetch_one(cur, f"SELECT {stat} FROM {q(table)}")}]
    return {"column": column, "by": by, "boxes": [{k: _json(v) for k, v in r.items()} for r in rows if r.get("n")]}


def scatter(table: str, x: str, y: str, limit: int = 1000) -> dict:
    cols = table_columns(table)
    xc, yc = _col(cols, x, ("number",), "x column"), _col(cols, y, ("number",), "y column")
    limit = max(10, min(limit, 5000))
    with get_cursor() as cur:
        total = cur.execute(f"SELECT COUNT(*) FROM {q(table)} WHERE {q(xc['name'])} IS NOT NULL AND {q(yc['name'])} IS NOT NULL").fetchone()[0]
        rows = fetch_all(cur, f"SELECT CAST({q(xc['name'])} AS DOUBLE) AS x, CAST({q(yc['name'])} AS DOUBLE) AS y FROM {q(table)} WHERE {q(xc['name'])} IS NOT NULL AND {q(yc['name'])} IS NOT NULL USING SAMPLE {limit} ROWS")
        corr = cur.execute(f"SELECT corr({q(xc['name'])}, {q(yc['name'])}) FROM {q(table)}").fetchone()[0]
    return {"x": x, "y": y, "n": total, "shown": len(rows), "correlation": _json(corr), "points": [{"x": _json(r["x"]), "y": _json(r["y"])} for r in rows]}
